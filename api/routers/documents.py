import csv
import io
import json
import os

from fastapi import APIRouter, BackgroundTasks, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, JSONResponse, Response
from sqlalchemy import select
from sqlalchemy.orm import Session

from api.schemas import FieldsUpdate, UploadResponse
from api.tasks import process_document
from config.settings import settings
from db.models import Correction, Document
from db.session import get_db
from models.tagger.postprocess import normalise_date, parse_amount
from pipeline.preprocess import load_pages
from storage.file_handler import save_upload
from validation.rules import validate

router = APIRouter(prefix="/documents", tags=["documents"])

NUMERIC = ("subtotal", "tax", "total")
TEXT = ("vendor", "invoice_number")
EDITABLE = NUMERIC + TEXT + ("date",)


def _get(db: Session, doc_id: str) -> Document:
    doc = db.get(Document, doc_id)
    if doc is None:
        raise HTTPException(404, "Document not found")
    return doc


def _summary(d: Document) -> dict:
    f = d.extraction.fields if d.extraction else {}
    return {"id": d.id, "name": d.original_name, "status": d.status, "error": d.error,
            "created_at": d.created_at.isoformat(), "finished_at": d.finished_at.isoformat() if d.finished_at else None,
            "document_type": d.extraction.document_type if d.extraction else None,
            "vendor": f.get("vendor"), "total": f.get("total")}


def _result(d: Document) -> dict | None:
    e = d.extraction
    if e is None:
        return None
    return {"document_type": e.document_type, "type_confidence": e.type_confidence, "fields": e.fields,
            "line_items": e.line_items, "field_confidences": e.confidences, "validation": e.validation,
            "words": e.words, "image_size": e.image_size, "inference_ms": e.inference_ms, **(e.extra or {})}


@router.post("", response_model=UploadResponse, status_code=202)
def upload(background: BackgroundTasks, file: UploadFile = File(...), db: Session = Depends(get_db)):
    path, size = save_upload(file)
    doc = Document(original_name=file.filename or "upload", file_path=path, mime_type=file.content_type or "", size_bytes=size)
    db.add(doc)
    db.commit()
    background.add_task(process_document, doc.id)
    return UploadResponse(document_id=doc.id, status=doc.status)


@router.get("")
def list_documents(limit: int = Query(100, le=500), db: Session = Depends(get_db)):
    docs = db.scalars(select(Document).order_by(Document.created_at.desc()).limit(limit)).all()
    return [_summary(d) for d in docs]


@router.get("/{doc_id}")
def get_document(doc_id: str, db: Session = Depends(get_db)):
    d = _get(db, doc_id)
    return {"document": _summary(d), "result": _result(d)}


@router.delete("/{doc_id}", status_code=204)
def delete_document(doc_id: str, db: Session = Depends(get_db)):
    d = _get(db, doc_id)
    for p in (d.file_path, os.path.join(settings.upload_dir, f"{d.id}_p1.png")):
        if os.path.exists(p):
            os.remove(p)
    db.delete(d)
    db.commit()


@router.get("/{doc_id}/image")
def get_image(doc_id: str, db: Session = Depends(get_db)):
    """The page image the boxes refer to (first page; PDFs are rendered once and cached)."""
    d = _get(db, doc_id)
    if not os.path.exists(d.file_path):
        raise HTTPException(404, "File missing on disk")
    if d.file_path.lower().endswith(".pdf"):
        png = os.path.join(settings.upload_dir, f"{d.id}_p1.png")
        if not os.path.exists(png):
            load_pages(d.file_path, settings.pdf_dpi, poppler_path=settings.poppler_path)[0].save(png)
        return FileResponse(png, media_type="image/png")
    return FileResponse(d.file_path)


def _clean(field: str, value):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if field in NUMERIC:
        v = parse_amount(value) if isinstance(value, str) else float(value)
        if v is None:
            raise HTTPException(422, f"'{field}' must be a number")
        return v
    if field == "date":
        v = normalise_date(str(value))
        if v is None:
            raise HTTPException(422, "'date' must be a valid date (e.g. 2024-03-15 or 15/03/2024)")
        return v
    return str(value).strip()


@router.put("/{doc_id}/fields")
def save_corrections(doc_id: str, body: FieldsUpdate, db: Session = Depends(get_db)):
    d = _get(db, doc_id)
    e = d.extraction
    if e is None:
        raise HTTPException(409, "Document has no extraction yet")
    unknown = [f for f in body.fields if f not in EDITABLE]
    if unknown:
        raise HTTPException(422, f"Unknown field(s): {unknown}. Editable: {list(EDITABLE)}")

    fields, conf = dict(e.fields), dict(e.confidences)
    for name, raw in body.fields.items():
        new, old = _clean(name, raw), fields.get(name)
        if new == old:
            continue
        db.add(Correction(extraction_id=e.id, field=name, original=None if old is None else str(old),
                          corrected="" if new is None else str(new)))
        if new is None:
            fields.pop(name, None); conf.pop(name, None)
        else:
            fields[name], conf[name] = new, 1.0                        # human-verified
    if body.line_items is not None:
        items = [i.model_dump() for i in body.line_items]
        if items != e.line_items:
            db.add(Correction(extraction_id=e.id, field="line_items", original=json.dumps(e.line_items),
                              corrected=json.dumps(items)))
            e.line_items = items
    e.fields, e.confidences = fields, conf                              # re-assign so SQLAlchemy sees the change
    e.validation = validate(fields, e.line_items, conf, settings.review_conf_threshold,
                            e.document_type, e.type_confidence, settings.type_warn_conf)
    d.status = "NEEDS_REVIEW" if e.validation["needs_review"] else "COMPLETE"
    db.commit()
    return {"document": _summary(d), "result": _result(d)}


@router.get("/{doc_id}/export")
def export(doc_id: str, fmt: str = Query("json", pattern="^(json|csv|xlsx)$"), db: Session = Depends(get_db)):
    d = _get(db, doc_id)
    if d.extraction is None:
        raise HTTPException(409, "Nothing to export yet")
    f, items = d.extraction.fields, d.extraction.line_items or []
    base = os.path.splitext(d.original_name)[0] or "document"
    head = ["vendor", "date", "invoice_number", "subtotal", "tax", "total"]
    if fmt == "json":
        return JSONResponse({"document_type": d.extraction.document_type, "fields": f, "line_items": items,
                             "validation": d.extraction.validation},
                            headers={"Content-Disposition": f'attachment; filename="{base}.json"'})
    if fmt == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(head + ["item_name", "item_qty", "item_price"])
        for it in items or [{}]:
            w.writerow([f.get(k, "") for k in head] + [it.get("name", ""), it.get("qty", ""), it.get("price", "")])
        return Response(buf.getvalue(), media_type="text/csv",
                        headers={"Content-Disposition": f'attachment; filename="{base}.csv"'})
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "fields"
    ws.append(["field", "value"])
    for k in head:
        ws.append([k, f.get(k)])
    ws2 = wb.create_sheet("line_items")
    ws2.append(["name", "qty", "price"])
    for it in items:
        ws2.append([it.get("name"), it.get("qty"), it.get("price")])
    out = io.BytesIO()
    wb.save(out)
    return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{base}.xlsx"'})
